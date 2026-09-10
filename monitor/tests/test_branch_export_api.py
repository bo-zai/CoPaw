"""HTTP export contract, including validation before database work."""

from io import BytesIO
from unittest.mock import AsyncMock
from urllib.parse import unquote

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from monitor.app.models.cron import CronBranchRankingResponse
from monitor.app.routers import api_router as router
from monitor.app.routers.cron_branch_export import XLSX_MEDIA_TYPE
from monitor.app.services.cron import get_query_service

URL = "/api/monitor/cron/export-branch-dimension"
DATES = {"start_date": "2026-09-01", "end_date": "2026-09-09"}


@pytest.fixture
def api():
    service = AsyncMock()
    service.get_branch_behavior.return_value = CronBranchRankingResponse(
        **DATES,
        items=[],
    )
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_query_service] = lambda: service
    with TestClient(app) as client:
        yield client, service


def test_binary_response_and_filters(api):
    client, service = api
    response = client.get(
        URL,
        params={
            **DATES,
            "bbk_ids": "100, 200,100",
            "sort_by": "totalTasks",
            "sort_order": "asc",
        },
        headers={"X-Source-Id": "source-a"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert unquote(response.headers["content-disposition"]).startswith(
        "attachment; filename*=UTF-8''定时任务分行维度_",
    )
    assert response.content.startswith(b"PK")
    assert load_workbook(BytesIO(response.content)).active.max_column == 22
    service.get_branch_behavior.assert_awaited_once_with(
        **DATES,
        bbk_ids="100,200",
        source_id="source-a",
    )


@pytest.mark.parametrize(
    "params",
    [
        {},
        {"start_date": "2026-09-01"},
        {**DATES, "start_date": "2026-02-30"},
        {**DATES, "start_date": "2026-9-01"},
        {**DATES, "end_date": "2026-08-01"},
        {**DATES, "sort_by": "totalTasks"},
        {**DATES, "sort_order": "asc"},
        {**DATES, "sort_by": "totalTasks", "sort_order": "up"},
        {**DATES, "sort_by": "total_tasks; DROP TABLE x", "sort_order": "asc"},
        {**DATES, "bbk_ids": "100,,200"},
        {**DATES, "bbk_ids": ""},
    ],
)
def test_invalid_params_are_chinese_errors_without_queries(api, params):
    client, service = api
    response = client.get(URL, params=params)
    assert response.status_code == 400
    assert isinstance(response.json()["detail"], str)
    service.get_branch_behavior.assert_not_awaited()


@pytest.mark.parametrize("status", [401, 403, 500])
def test_errors_are_non_200_json(api, status):
    client, service = api
    service.get_branch_behavior.side_effect = (
        HTTPException(status, "无权访问")
        if status != 500
        else RuntimeError("private database connection information")
    )
    response = client.get(URL, params=DATES)
    assert response.status_code == status
    assert isinstance(response.json()["detail"], str)
    assert "private" not in response.text
    assert "content-disposition" not in response.headers
